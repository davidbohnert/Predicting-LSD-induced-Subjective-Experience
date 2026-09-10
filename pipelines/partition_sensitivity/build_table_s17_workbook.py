"""Write Table S17 as flat CSVs and a compact XLSX.

Same numbers as the supplement tables, in a form that can be sorted and
checked.  The workbook carries the columns the document leaves out --
the raw fold-selection count, the anatomical-overlap percentage behind
each ROI descriptor, and the edge index into the saved frequency vectors.
"""
from __future__ import annotations

import csv
import argparse
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import table_s17_data as data

OUTPUT = Path(__file__).resolve().parent / "output"
HEADER_FILL = PatternFill("solid", fgColor="F3F3F3")

PANEL_A_HEADER = [
    "Outcome", "Rank",
    "ROI 1 ID", "ROI 1 parcel", "ROI 1 anatomy", "ROI 1 overlap %",
    "ROI 1 x", "ROI 1 y", "ROI 1 z", "ROI 1 network",
    "ROI 2 ID", "ROI 2 parcel", "ROI 2 anatomy", "ROI 2 overlap %",
    "ROI 2 x", "ROI 2 y", "ROI 2 z", "ROI 2 network",
    "Network pair", "Fixed-partition folds", "Consensus recurrence",
    "Fold selection count", "Selection frequency %", "Mean partial r",
    "Maximally stable edges in outcome", "Edge index",
]
PANEL_B_HEADER = [
    "Outcome", "Rank", "ROI ID", "Parcel", "Anatomy", "Overlap %",
    "x", "y", "z", "Network", "Fixed-mask degree",
    "Stability-weighted degree", "Edge-endpoint share %",
]


def panel_a_records(rows):
    for row in rows:
        first, second = row["ROI1"], row["ROI2"]
        yield [
            row["Outcome"], row["Rank"],
            first.parcel_id, first.label, first.anatomy, first.anatomy_overlap,
            first.x, first.y, first.z, first.network,
            second.parcel_id, second.label, second.anatomy, second.anatomy_overlap,
            second.x, second.y, second.z, second.network,
            row["NetworkPair"].replace("–", "-"),
            f'{row["FixedPartitionFolds"]}/{row["NoFolds"]}',
            f'{row["ConsensusRecurrence"]}/{row["NoPartitions"]}',
            row["FoldSelectionCount"],
            round(row["SelectionFrequency"], 1),
            round(row["MeanPartialR"], 3),
            row["MaximallyStableEdges"], row["EdgeIndex"],
        ]


def panel_b_records(rows):
    for row in rows:
        parcel = row["Parcel"]
        yield [
            row["Outcome"], row["Rank"], parcel.parcel_id, parcel.label,
            parcel.anatomy, parcel.anatomy_overlap,
            parcel.x, parcel.y, parcel.z, row["Network"],
            row["FixedMaskDegree"], round(row["WeightedDegree"], 1),
            round(row["EndpointShare"], 2),
        ]


def write_csv(path, header, records):
    with path.open("w", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(header)
        writer.writerows(records)


def add_sheet(book, title, header, records):
    sheet = book.create_sheet(title)
    sheet.append(header)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for record in records:
        sheet.append(record)
    sheet.freeze_panes = "A2"
    for index, name in enumerate(header, start=1):
        widest = max([len(name)] + [len(str(sheet.cell(r, index).value or ""))
                                    for r in range(2, sheet.max_row + 1)])
        sheet.column_dimensions[get_column_letter(index)].width = min(max(widest + 2, 9), 34)
    return sheet


def add_parameters_sheet(book, stability):
    sheet = book.create_sheet("Parameters")
    reference = stability["GDE"]
    rows = [
        ("Connectome", "LSD_difference (LSD-placebo difference)"),
        ("Covariates", "covars_indicator.mat [study_2, study_3, sex, age, mean_FD]"),
        ("Network direction", "positive"),
        ("Cross-validation partitions", f'{reference["no_partitions"]} (CV seeds 123-222, as in Figure S1)'),
        ("Folds per partition", reference["no_folds"]),
        ("Training folds in total", reference["no_training_folds"]),
        ("Edge-selection threshold", "p < .01"),
        ("Consensus rule", f'selected in at least {reference["consensus_folds_required"]} of {reference["no_folds"]} folds'),
        ("Fixed partition", "CV seed 123, the partition used throughout the manuscript"),
        ("Correlation method", "; ".join(
            f'{o} partial {stability[o]["correlation"]} (n = {stability[o]["n"]})'
            for o in data.OUTCOMES)),
        ("Consensus edges, fixed partition", "; ".join(
            f'{o} {stability[o]["fixed_edges"]}' for o in data.OUTCOMES)),
        ("Consensus recurrence", "partitions in which the edge met the consensus rule"),
        ("Selection frequency", "training folds selecting the edge, as a percentage of all 1,000"),
        ("Mean partial r", "mean edge-behaviour partial correlation across all 1,000 training folds"),
        ("Stability-weighted degree",
         "sum, over the 415 edges at a node, of the proportion of the 1,000 folds selecting each edge"),
        ("Edge-endpoint share", "a node's stability-weighted degree over the sum across all 416 nodes"),
        ("Anatomy", "AAL region with the largest voxel overlap with the parcel; overlap % is given"),
        ("Source", "Generated by edge_node_stability_main.m"),
    ]
    for key, value in rows:
        sheet.append([key, value])
    for cell in sheet["A"]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(vertical="top")
    for cell in sheet["B"]:
        cell.alignment = Alignment(vertical="top", wrap_text=True)
    sheet.column_dimensions["A"].width = 34
    sheet.column_dimensions["B"].width = 78


def main() -> None:
    parser = argparse.ArgumentParser(description="Build Table S17 source files")
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    global OUTPUT
    OUTPUT = args.output_dir.resolve()
    OUTPUT.mkdir(parents=True, exist_ok=True)
    data.OUTPUT = OUTPUT
    xlsx = OUTPUT / "table_S17_edge_and_node_stability.xlsx"
    panel_a, panel_b, stability = data.load_panels()
    write_csv(OUTPUT / "table_S17_panel_A_top_edges.csv",
              PANEL_A_HEADER, panel_a_records(panel_a))
    write_csv(OUTPUT / "table_S17_panel_B_top_nodes.csv",
              PANEL_B_HEADER, panel_b_records(panel_b))

    book = Workbook()
    book.remove(book.active)
    add_sheet(book, "Panel A edges", PANEL_A_HEADER, list(panel_a_records(panel_a)))
    add_sheet(book, "Panel B nodes", PANEL_B_HEADER, list(panel_b_records(panel_b)))
    add_parameters_sheet(book, stability)
    book.save(xlsx)
    print(f"Panel A rows: {len(panel_a)}")
    print(f"Panel B rows: {len(panel_b)}")
    print(f"Wrote {xlsx}")


if __name__ == "__main__":
    main()
